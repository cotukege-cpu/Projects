import tkinter as tk
from tkinter import messagebox, ttk
import mysql.connector
import matplotlib.pyplot as plt
import seaborn as sns
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import pandas as pd
import openpyxl
from openpyxl import load_workbook, Workbook
import os
import subprocess
import sys

EXCEL_FILE = "foote_changed.xlsx"



def get_db_connection():
    return mysql.connector.connect(
        host=
        user=
        password=
        database=
    )


def append_player_to_excel(player):
    headers = ["Name", "Age", "Position", "Goals", "Performance", "Team", "League"]

    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        # Eğer dosya boşsa header ekle
        if ws.max_row == 0 or (ws.max_row == 1 and ws.cell(1, 1).value is None):
            ws.append(headers)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Players"
        ws.append(headers)

    ws.append([
        player["Name"],
        int(player["Age"]),
        player["Position"],
        int(player["Goals"]),
        float(player["Performance"]),
        player["Team"],
        player["League"]
    ])
    wb.save(EXCEL_FILE)


def open_excel_file():
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT MIN(id), Player_name, Age, Position, Goals, Performance, Team, League FROM foote GROUP BY Player_name, Age, Position, Goals, Performance, Team, League")
        rows = cursor.fetchall()
        cursor.close()
        conn.close()

        wb = Workbook()
        ws = wb.active
        ws.title = "Players"
        ws.append(["Name", "Age", "Position", "Goals", "Performance", "Team", "League"])

        for row in rows:
            _, name, age, position, goals, perf, team, league = row
            ws.append([name, age, position, goals, perf, team, league])

        wb.save(EXCEL_FILE)

        if sys.platform.startswith("win"):
            os.startfile(os.path.abspath(EXCEL_FILE))
        elif sys.platform.startswith("darwin"):
            subprocess.call(["open", os.path.abspath(EXCEL_FILE)])
        else:
            subprocess.call(["xdg-open", os.path.abspath(EXCEL_FILE)])

    except Exception as e:
        messagebox.showerror("Error", f"File could not be opened:\n{e}")


def login_screen():
    for widget in root.winfo_children():
        widget.destroy()

    root.configure(bg="lightblue")

    title = tk.Label(root, text="FOOTBALL ANALYSING SYSTEM",
                     font=("Arial", 40, "bold"), bg="lightblue")
    title.pack(pady=50)

    tk.Label(root, text="Username", bg="lightblue", font=("Arial", 14)).pack()
    username_entry = tk.Entry(root, font=("Arial", 14))
    username_entry.pack(pady=10)

    tk.Label(root, text="Password", bg="lightblue", font=("Arial", 14)).pack()
    password_entry = tk.Entry(root, show="*", font=("Arial", 14))
    password_entry.pack(pady=10)

    def login():
        username = username_entry.get()
        password = password_entry.get()

        if username == "scout" and password == "scout123":
            scout_screen()
        elif username == "manager" and password == "manager123":
            manager_screen()
        else:
            messagebox.showerror("Error", "TRY AGAIN!!!")

    tk.Button(root, text="Login", command=login, font=("Arial", 14)).pack(pady=20)

    ball = tk.Label(root, text="⚽", font=("Arial", 90), bg="lightblue")
    ball.place(x=50, y=200)


def scout_screen():
    for widget in root.winfo_children():
        widget.destroy()

    root.configure(bg="lightblue")

    title = tk.Label(root, text="SCOUT PANEL",
                     font=("Arial", 30, "bold"), bg="lightblue")
    title.pack(pady=20)

    league_teams = {
        "Premier League": ["Arsenal", "Chelsea", "Liverpool",
                           "Manchester City", "Manchester United", "Tottenham"],
        "Bundesliga": ["Bayern Munich", "Borussia Dortmund"],
        "Serie A": ["Milan", "Inter Milan", "Juventus"],
        "Ligue 1": ["Paris Saint-Germain"],
        "LaLiga": ["Atletico Madrid", "Barcelona", "Real Madrid"]
    }

    entries = {}
    fields = ["Name", "Age", "Position", "Goals", "Performance", "Team", "League"]
    for field in fields:
        tk.Label(root, text=field, bg="lightblue", font=("Arial", 14)).pack()
        if field == "League":
            leagues = ["Premier League", "Bundesliga", "Serie A", "Ligue 1", "LaLiga"]
            leagues_var = tk.StringVar()
            combo = ttk.Combobox(root, textvariable=leagues_var, values=leagues, state='readonly', font=("Arial", 14))
            combo.pack(pady=5)
            entries[field] = leagues_var

            def on_league_change(event):
                selected_league = leagues_var.get()
                teams = league_teams.get(selected_league, [])
                entries["Team"].set("")
                entries["Team"].widget.configure(values=teams, state='readonly')

            combo.bind("<<ComboboxSelected>>", on_league_change)

        elif field == "Position":
            positions = ["ST", "RW", "LW", "CAM", "CM", "CDM", "RB", "CB", "LB", "GK"]
            position_var = tk.StringVar()
            combo = ttk.Combobox(root, textvariable=position_var, values=positions, state='readonly',
                                 font=("Arial", 14))
            combo.pack(pady=5)
            entries[field] = position_var

        elif field == "Team":
            team_var = tk.StringVar()
            team_combo = ttk.Combobox(root, textvariable=team_var, values=[], state='disabled', font=("Arial", 14))
            team_combo.pack(pady=5)
            team_var.widget = team_combo
            entries[field] = team_var

        else:
            entry = tk.Entry(root, font=("Arial", 16))
            entry.pack(pady=5)
            entries[field] = entry

    def save_player():
        player = {}

        try:
            age = int(entries["Age"].get())
            goals = int(entries["Goals"].get())
            perf = float(entries["Performance"].get())
        except:
            messagebox.showerror("Error", "Please enter valid numbers!")
            return

        if age < 0 or goals < 0 or perf < 0:
            messagebox.showerror("Error", "Negative values are not allowed!")
            return

        if perf > 10:
            messagebox.showerror("Error", "You can enter maximum 10 for performance!")
            return

        for field in fields:
            player[field] = entries[field].get()

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
                       INSERT INTO foote (Player_name, Age, Position, Goals, Performance, Team, League)
                       VALUES (%s, %s, %s, %s, %s, %s, %s)
                       """, (
                           player["Name"], int(player["Age"]), player["Position"],
                           int(player["Goals"]), float(player["Performance"]),
                           player["Team"], player["League"]
                       ))
        conn.commit()
        cursor.close()
        conn.close()

        try:
            append_player_to_excel(player)
        except Exception as e:
            messagebox.showwarning("Warning!", f"Couldnt add the player:\n{e}")

        messagebox.showinfo("Saved", "Player added!")

        for field in fields:
            if field == "Team":
                entries[field].set("")
                entries[field].widget.configure(values=[], state='disabled')
            elif field == "League":
                entries[field].set("")
            elif field == "Position":
                entries[field].set("")
            else:
                entries[field].delete(0, tk.END)

    tk.Button(root, text="Save Player", command=save_player,
              font=("Arial", 14)).pack(pady=20)

    tk.Button(root, text="Logout", command=login_screen,
              font=("Arial", 12)).pack()

    ball = tk.Label(root, text="⚽", font=("Arial", 90), bg="lightblue")
    ball.place(x=50, y=200)


def manager_screen():
    for widget in root.winfo_children():
        widget.destroy()

    root.configure(bg="lightblue")

    title = tk.Label(root, text="MANAGER PANEL",
                     font=("Arial", 30, "bold"), bg="lightblue")
    title.pack(pady=20)

    frame = tk.Frame(root, bg="lightblue")
    frame.pack()

    tk.Label(frame, text="Name", bg="lightblue").grid(row=0, column=0)
    name_entry = tk.Entry(frame)
    name_entry.grid(row=0, column=1)

    tk.Label(frame, text="Age Min", bg="lightblue").grid(row=1, column=0)
    age_min = tk.Entry(frame)
    age_min.grid(row=1, column=1)

    tk.Label(frame, text="Age Max", bg="lightblue").grid(row=1, column=2)
    age_max = tk.Entry(frame)
    age_max.grid(row=1, column=3)

    tk.Label(frame, text="Position", bg="lightblue").grid(row=2, column=0)
    positions = ["ST", "RW", "LW", "CAM", "CM", "CDM", "RB", "CB", "LB", "GK"]
    position_var = tk.StringVar()
    position_combo = ttk.Combobox(frame, textvariable=position_var, values=positions, state='readonly',
                                  font=("Arial", 7))
    position_combo.grid(row=2, column=1)

    tk.Label(frame, text="Goals Min", bg="lightblue").grid(row=3, column=0)
    goals_min = tk.Entry(frame)
    goals_min.grid(row=3, column=1)

    tk.Label(frame, text="Goals Max", bg="lightblue").grid(row=3, column=2)
    goals_max = tk.Entry(frame)
    goals_max.grid(row=3, column=3)

    tk.Label(frame, text="Performance Min", bg="lightblue").grid(row=4, column=0)
    perf_min = tk.Entry(frame)
    perf_min.grid(row=4, column=1)

    tk.Label(frame, text="Performance Max", bg="lightblue").grid(row=4, column=2)
    perf_max = tk.Entry(frame)
    perf_max.grid(row=4, column=3)

    tk.Label(frame, text="Team", bg="lightblue").grid(row=5, column=0)
    team_entry = tk.Entry(frame)
    team_entry.grid(row=5, column=1)

    tk.Label(frame, text="League", bg="lightblue").grid(row=5, column=2)
    leagues = ["Premier League", "Bundesliga", "Serie A", "Ligue 1", "LaLiga"]
    leagues_var = tk.StringVar()
    l_combo = ttk.Combobox(frame, textvariable=leagues_var, values=leagues, state='readonly', font=("Arial", 7))
    l_combo.grid(row=5, column=3)

    columns = ("Name", "Team", "League", "Age", "Position", "Goals", "Performance")
    tree = ttk.Treeview(root, columns=columns, show="headings", height=15)

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=150, anchor="center")

    tree.pack(pady=20)

    def filter_players():
        tree.delete(*tree.get_children())

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT MIN(id), Player_name, Team, League, Age, Position, Goals, Performance FROM foote GROUP BY Player_name, Team, League, Age, Position, Goals, Performance")
        db_players = cursor.fetchall()
        cursor.close()
        conn.close()

        for row in db_players:
            _, name, team, league, age, position, goals, perf = row

            if name_entry.get() and name_entry.get().lower() not in name.lower():
                continue
            if position_var.get() and position_var.get() != position:
                continue
            if team_entry.get() and team_entry.get().lower() not in team.lower():
                continue
            if leagues_var.get() and leagues_var.get() != league:
                continue
            if age_min.get() and age < int(age_min.get()):
                continue
            if age_max.get() and age > int(age_max.get()):
                continue
            if goals_min.get() and goals < int(goals_min.get()):
                continue
            if goals_max.get() and goals > int(goals_max.get()):
                continue
            if perf_min.get() and perf < float(perf_min.get()):
                continue
            if perf_max.get() and perf > float(perf_max.get()):
                continue

            tree.insert("", tk.END, values=(name, team, league, age, position, goals, perf))

    def delete_player():
        selected = tree.selection()

        if not selected:
            messagebox.showwarning("Warning", "Select a player!")
            return

        index = selected[0]
        selected_values = tree.item(index, "values")

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
                       DELETE
                       FROM foote
                       WHERE Player_name = %s
                         AND Team = %s
                         AND League = %s
                       """, (selected_values[0], selected_values[1], selected_values[2]))
        conn.commit()
        cursor.close()
        conn.close()

        tree.delete(index)
        messagebox.showinfo("Deleted", "Player removed!")

    btn_row = tk.Frame(root, bg="lightblue")
    btn_row.pack(pady=10)

    tk.Button(btn_row, text="Filter", command=filter_players,
              font=("Arial", 12), width=15).pack(side="left", padx=10)

    tk.Button(btn_row, text="Delete Selected", command=delete_player,
              bg="red", fg="white", font=("Arial", 12), width=15).pack(side="left", padx=10)

    tk.Button(btn_row, text="⚽ Dream 11", command=dream11_screen,
              bg="green", fg="white", font=("Arial", 12, "bold"), width=15).pack(side="left", padx=10)

    tk.Button(btn_row, text="📊 Graphics", command=graphics_screen,
              bg="darkblue", fg="white", font=("Arial", 12, "bold"), width=15).pack(side="left", padx=10)

    tk.Button(btn_row, text="📂 Open Excel", command=open_excel_file,
              bg="darkorange", fg="white", font=("Arial", 12, "bold"), width=15).pack(side="left", padx=10)

    tk.Button(root, text="Logout", command=login_screen).pack(pady=10)

    tk.Label(root, text="⚽", font=("Arial", 90), bg="lightblue").place(x=50, y=200)

    filter_players()


def dream11_screen():
    window = tk.Toplevel(root)
    window.title("Dream 11")
    window.configure(bg="lightblue")

    screen_w = root.winfo_screenwidth()
    screen_h = root.winfo_screenheight()
    win_w = int(screen_w * 0.85)
    win_h = int(screen_h * 0.85)
    x = (screen_w - win_w) // 2
    y = (screen_h - win_h) // 2
    window.geometry(f"{win_w}x{win_h}+{x}+{y}")

    tk.Label(window, text="DREAM 11 GENERATOR", font=("Arial", 24, "bold"), bg="lightblue").pack(pady=15)

    container = tk.Frame(window, bg="lightblue")
    container.pack(fill="both", expand=True, padx=20, pady=10)

    container.columnconfigure(0, weight=1)
    container.columnconfigure(1, weight=0)
    container.columnconfigure(2, weight=1)
    container.rowconfigure(0, weight=1)

    left_frame = tk.Frame(container, bg="lightblue", bd=2, relief="groove")
    left_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))

    tk.Label(left_frame, text="📋 All Players", font=("Arial", 14, "bold"), bg="lightblue").pack(pady=10)

    left_cols = ("Name", "Team", "Position", "Goals", "Performance")
    left_tree = ttk.Treeview(left_frame, columns=left_cols, show="headings", height=20)
    for col in left_cols:
        left_tree.heading(col, text=col)
        left_tree.column(col, width=110, anchor="center")
    left_tree.pack(fill="both", expand=True, padx=10, pady=5)

    mid_frame = tk.Frame(container, bg="lightblue")
    mid_frame.grid(row=0, column=1, sticky="ns", padx=10)

    def add_to_dream11():
        selected = left_tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Select a player to add!", parent=window)
            return

        if len(right_tree.get_children()) >= 11:
            messagebox.showwarning("Warning", "Dream 11 is full! Max 11 players.", parent=window)
            return

        item = selected[0]
        values = left_tree.item(item, "values")

        for child in right_tree.get_children():
            if right_tree.item(child, "values")[0] == values[0]:
                messagebox.showwarning("Warning", "Player already in Dream 11!", parent=window)
                return

        right_tree.insert("", tk.END, values=values)
        update_count()

    def remove_from_dream11():
        selected = right_tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Select a player to remove!", parent=window)
            return
        right_tree.delete(selected[0])
        update_count()

    def update_count():
        count = len(right_tree.get_children())
        count_label.config(text=f"Players: {count} / 11",
                           fg="green" if count == 11 else "red")

    tk.Button(mid_frame, text="➕ Add Player", command=add_to_dream11,
              bg="green", fg="white", font=("Arial", 13, "bold"),
              width=15, pady=8).pack(pady=10)

    tk.Button(mid_frame, text="➖ Remove Player", command=remove_from_dream11,
              bg="red", fg="white", font=("Arial", 13, "bold"),
              width=15, pady=8).pack(pady=10)

    count_label = tk.Label(mid_frame, text="Players: 0 / 11",
                           font=("Arial", 12, "bold"), bg="lightblue", fg="red")
    count_label.pack(pady=5)

    right_frame = tk.Frame(container, bg="lightblue", bd=2, relief="groove")
    right_frame.grid(row=0, column=2, sticky="nsew", padx=(10, 0))

    tk.Label(right_frame, text="⭐ Dream 11 Squad", font=("Arial", 14, "bold"), bg="lightblue").pack(pady=10)

    right_cols = ("Name", "Team", "Position", "Goals", "Performance")
    right_tree = ttk.Treeview(right_frame, columns=right_cols, show="headings", height=20)
    for col in right_cols:
        right_tree.heading(col, text=col)
        right_tree.column(col, width=110, anchor="center")
    right_tree.pack(fill="both", expand=True, padx=10, pady=5)

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT MIN(id), Player_name, Team, Position, Goals, Performance FROM foote GROUP BY Player_name, Team, Position, Goals, Performance")
    all_players = cursor.fetchall()
    cursor.close()
    conn.close()

    for row in all_players:
        _, name, team, position, goals, perf = row
        left_tree.insert("", tk.END, values=(name, team, position, goals, perf))


login_screen()
root.mainloop()


def graphics_screen():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT MIN(id), Player_name, Team, League, Age, Position, Goals, Performance FROM foote GROUP BY Player_name, Team, League, Age, Position, Goals, Performance")
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    if not rows:
        messagebox.showwarning("Warning", "No data found in database!")
        return

    df = pd.DataFrame(rows, columns=["id", "Name", "Team", "League", "Age", "Position", "Goals", "Performance"])
    df.drop(columns=["id"], inplace=True)
    df["Age"] = pd.to_numeric(df["Age"])
    df["Goals"] = pd.to_numeric(df["Goals"])
    df["Performance"] = pd.to_numeric(df["Performance"])

    window = tk.Toplevel(root)
    window.title("Graphics Dashboard")
    window.configure(bg="lightblue")

    screen_w = root.winfo_screenwidth()
    screen_h = root.winfo_screenheight()
    win_w = int(screen_w * 0.95)
    win_h = int(screen_h * 0.90)
    x = (screen_w - win_w) // 2
    y = (screen_h - win_h) // 2
    window.geometry(f"{win_w}x{win_h}+{x}+{y}")

    tk.Label(window, text="📊 GRAPHICS DASHBOARD", font=("Arial", 20, "bold"), bg="lightblue").pack(pady=10)

    notebook = ttk.Notebook(window)
    notebook.pack(fill="both", expand=True, padx=10, pady=10)

    def embed_chart(fig, parent):
        canvas = FigureCanvasTkAgg(fig, master=parent)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True)

    tab1 = tk.Frame(notebook, bg="white")
    notebook.add(tab1, text="⚽ Goals by Position")
    fig1, ax1 = plt.subplots(figsize=(10, 5))
    sns.barplot(data=df, x="Position", y="Goals", palette="viridis", ax=ax1)
    ax1.set_title("Average Goals by Position", fontsize=14, fontweight="bold")
    ax1.set_xlabel("Position")
    ax1.set_ylabel("Average Goals")
    plt.tight_layout()
    embed_chart(fig1, tab1)

    tab2 = tk.Frame(notebook, bg="white")
    notebook.add(tab2, text="📈 Performance Distribution")
    fig2, ax2 = plt.subplots(figsize=(10, 5))
    sns.histplot(data=df, x="Performance", bins=10, kde=True, color="steelblue", ax=ax2)
    ax2.set_title("Performance Distribution of All Players", fontsize=14, fontweight="bold")
    ax2.set_xlabel("Performance Score")
    ax2.set_ylabel("Number of Players")
    plt.tight_layout()
    embed_chart(fig2, tab2)

    tab3 = tk.Frame(notebook, bg="white")
    notebook.add(tab3, text="🎯 Goals vs Performance")
    fig3, ax3 = plt.subplots(figsize=(10, 5))
    sns.scatterplot(data=df, x="Goals", y="Performance", hue="Position", palette="tab10", s=100, ax=ax3)
    ax3.set_title("Goals vs Performance by Position", fontsize=14, fontweight="bold")
    ax3.set_xlabel("Goals")
    ax3.set_ylabel("Performance")
    plt.tight_layout()
    embed_chart(fig3, tab3)

    tab4 = tk.Frame(notebook, bg="white")
    notebook.add(tab4, text="🏆 Players by League")
    fig4, ax4 = plt.subplots(figsize=(10, 5))
    sns.countplot(data=df, x="League", palette="Set2", ax=ax4)
    ax4.set_title("Number of Players per League", fontsize=14, fontweight="bold")
    ax4.set_xlabel("League")
    ax4.set_ylabel("Number of Players")
    plt.xticks(rotation=15)
    plt.tight_layout()
    embed_chart(fig4, tab4)

    tab5 = tk.Frame(notebook, bg="white")
    notebook.add(tab5, text="🔢 Age vs Goals")
    fig5, ax5 = plt.subplots(figsize=(10, 5))
    sns.regplot(data=df, x="Age", y="Goals", scatter_kws={"color": "coral", "s": 80},
                line_kws={"color": "darkred"}, ax=ax5)
    ax5.set_title("Age vs Goals (with Trend Line)", fontsize=14, fontweight="bold")
    ax5.set_xlabel("Age")
    ax5.set_ylabel("Goals")
    plt.tight_layout()
    embed_chart(fig5, tab5)
